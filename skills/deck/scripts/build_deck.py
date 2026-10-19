#!/usr/bin/env python3
"""Build a consulting .pptx (white background, grey boxes, one accent colour) from a storyboard JSON.

Usage:
    python3 build_deck.py storyboard.json -o deck.pptx

The storyboard format is described in references/storyboard-schema.md.
Layout warnings (text too long for its box) go to stderr.
"""
import argparse
import json
import math
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from layouts_spec import HARVEY, LAYOUTS  # noqa: E402

# ---------------------------------------------------------------- house style
SW, SH = 13.333, 7.5
ML, MR = 0.5, 0.5
CX, CY = ML, 1.55
CW, CH = SW - ML - MR, 5.25  # content area: 1.55 -> 6.80 in

BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREY_TEXT = RGBColor(0x59, 0x59, 0x59)
GREY_BOX = RGBColor(0xF2, 0xF2, 0xF2)
GREY_MID = RGBColor(0xA6, 0xA6, 0xA6)
GREY_DARK = RGBColor(0x7F, 0x7F, 0x7F)
GREY_LINE = RGBColor(0xBF, 0xBF, 0xBF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
CALC_BLUE = RGBColor(0x00, 0x70, 0xC0)  # values computed by the team, not read in a source
GROUP_FILLS = [RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x2E, 0x75, 0xB6), RGBColor(0x9D, 0xC3, 0xE6),
               RGBColor(0x59, 0x59, 0x59), RGBColor(0xA6, 0xA6, 0xA6), RGBColor(0xD9, 0xD9, 0xD9)]
GROUP_TEXT_DARK = {2, 4, 5}  # light fills take dark text
FONT = "Arial"

WARNINGS = []
_CTX = {"slide": 0}


def warn(msg):
    WARNINGS.append(f"slide {_CTX['slide']}: {msg}")


# ---------------------------------------------------------------- text
def _para(item):
    """Normalise a paragraph: str or dict {text, bold, bullet, color, rel_size, space_before, align, italic}."""
    return {"text": item} if isinstance(item, str) else item


def estimate_height(paras, width_in, size):
    """Estimated height in inches for Arial: average glyph width ≈ 0.5 × size (0.56 bold)."""
    total = 0.0
    for p in paras:
        s = size * p.get("rel_size", 1.0)
        w = width_in - (0.2 if p.get("bullet") else 0)
        cpl = max(1, int(w * 72 / (s * (0.56 if p.get("bold") else 0.5))))
        longest = max((len(x) for x in str(p["text"]).split()), default=0)
        if longest > cpl:  # a word would break in two: size too large
            return 99.0
        lines = sum(max(1, math.ceil(len(chunk) / cpl)) for chunk in str(p["text"]).split("\n"))
        total += lines * s * 1.18 / 72 + p.get("space_before", 3) / 72
    return total


def fit_size(paras, w, h, size, min_size, quiet=False):
    s = size
    while s > min_size and estimate_height(paras, w, s) > h:
        s -= 0.5
    if estimate_height(paras, w, s) > h and not quiet:
        warn(f"text too long for its box ({str(paras[0]['text'])[:50]}...): shorten")
    return s


def _set_bullet(p, char):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.18)))
    pPr.set("indent", str(-Inches(0.18)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def fill_text(tf, paras, w, h, size=12, min_size=9, color=BLACK, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP, margin=0.1, fit=True):
    paras = [_para(p) for p in paras if p is not None and _para(p).get("text", "") != ""]
    if not paras:
        return
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin * 0.8)
    s = fit_size(paras, w - 2 * margin, h - 1.6 * margin, size, min_size) if fit else size
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        if i > 0:
            para.space_before = Pt(p.get("space_before", 3))
        run = para.add_run()
        run.text = str(p["text"])
        f = run.font
        f.name = FONT
        f.size = Pt(round(s * p.get("rel_size", 1.0) * 2) / 2)
        f.bold = bool(p.get("bold", False))
        f.italic = bool(p.get("italic", False))
        f.color.rgb = p.get("color", color)
        if p.get("bullet"):
            _set_bullet(para, p.get("bullet_char", "•"))


def textbox(slide, x, y, w, h, paras, **kw):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(shape.text_frame, paras, w, h, **kw)
    return shape


def rect(slide, x, y, w, h, fill=GREY_BOX, line=None, kind=MSO_SHAPE.RECTANGLE, line_w=1.0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    sty = s._element.find(qn("p:style"))
    if sty is not None:  # the theme style adds a shadow and implicit colours
        s._element.remove(sty)
    return s


def block_paras(title, points, title_color=BLACK, text_color=BLACK):
    paras = []
    if title:
        paras.append({"text": title, "bold": True, "rel_size": 1.08, "color": title_color})
    for pt in points or []:
        paras.append({"text": pt, "bullet": True, "color": text_color, "space_before": 5})
    return paras


def block(slide, x, y, w, h, title, points, fill=GREY_BOX, line=None, size=14,
          title_color=BLACK, text_color=BLACK, anchor=MSO_ANCHOR.TOP, fit=True):
    s = rect(slide, x, y, w, h, fill=fill, line=line)
    fill_text(s.text_frame, block_paras(title, points, title_color, text_color), w, h,
              size=size, anchor=anchor, margin=0.15, fit=fit)
    return s


def common_size(items, size=14, min_size=9, margin=0.15):
    """One font size for sibling blocks: (w, h, title, points) -> largest size that fits everywhere.
    Neighbouring blocks in different sizes break layer-cake scanning."""
    best = size
    for w, h, title, points in items:
        paras = [_para(p) for p in block_paras(title, points)]
        if paras:
            best = min(best, fit_size(paras, w - 2 * margin, h - 1.6 * margin, size, min_size))
    return best


def line(slide, x1, y1, x2, y2, color=GREY_LINE, width=1.0, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    sty = c._element.find(qn("p:style"))
    if sty is not None:  # the theme style adds a shadow to connectors
        c._element.remove(sty)
    if arrow:
        etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")
    return c


def split_w(total, n, gap):
    w = (total - gap * (n - 1)) / n
    return [(i * (w + gap), w) for i in range(n)]


def side_blocks(slide, blocks, x, w, y=CY, h_total=CH):
    n = len(blocks)
    gap = 0.15
    h = (h_total - gap * (n - 1)) / n
    sz = common_size([(w, h, b.get("title"), b.get("points")) for b in blocks])
    for i, b in enumerate(blocks):
        block(slide, x, y + i * (h + gap), w, h, b.get("title"), b.get("points"), size=sz, fit=False,
              line=ACCENT if b.get("accent") else None, title_color=ACCENT if b.get("accent") else BLACK)


# ---------------------------------------------------------------- frame
def header(slide, s, num):
    textbox(slide, ML, 0.35, CW, 1.05, [{"text": s["title"], "bold": True}],
            size=22, min_size=15, margin=0.02)
    if s.get("source"):
        textbox(slide, ML, 6.92, CW - 1.0, 0.4, [f"Source: {s['source']}"], size=9, min_size=7,
                color=GREY_TEXT, margin=0.02, anchor=MSO_ANCHOR.BOTTOM)
    textbox(slide, SW - MR - 0.8, 6.92, 0.8, 0.4, [str(num)], size=9, color=GREY_TEXT,
            align=PP_ALIGN.RIGHT, margin=0.02, anchor=MSO_ANCHOR.BOTTOM, fit=False)


# ---------------------------------------------------------------- layouts
def l_cover(slide, s):
    rect(slide, 0, 0, 0.18, SH, fill=ACCENT)
    textbox(slide, 1.0, 2.3, SW - 2.0, 1.9, [{"text": s["title"], "bold": True}],
            size=34, min_size=24, anchor=MSO_ANCHOR.BOTTOM, margin=0.02)
    if s.get("context"):
        textbox(slide, 1.0, 4.35, SW - 2.0, 0.8, [s["context"]], size=16, color=GREY_TEXT, margin=0.02)


def l_executive_summary(slide, s):
    items = s["items"]
    n = len(items)
    gap = 0.12
    rh = (CH - gap * (n - 1)) / n
    nw = 0.55
    tw = CW - nw - 0.15
    paras_all = [[{"text": it["takeaway"], "bold": True}] +
                 [{"text": p, "bullet": True, "bullet_char": "–", "color": GREY_TEXT, "space_before": 3}
                  for p in it.get("points", [])] for it in items]
    sz = min(fit_size([_para(p) for p in paras], tw - 0.3, rh - 0.2, 14, 9) for paras in paras_all)
    for i, (it, paras) in enumerate(zip(items, paras_all)):
        y = CY + i * (rh + gap)
        num = rect(slide, CX, y, nw, rh, fill=ACCENT if it.get("accent") else GREY_DARK)
        fill_text(num.text_frame, [{"text": str(i + 1), "bold": True, "color": WHITE,
                                    "align": PP_ALIGN.CENTER}], nw, rh, size=16,
                  anchor=MSO_ANCHOR.MIDDLE, fit=False, margin=0.02)
        box = rect(slide, CX + nw + 0.15, y, tw, rh, fill=GREY_BOX)
        fill_text(box.text_frame, paras, tw, rh, size=sz, anchor=MSO_ANCHOR.MIDDLE,
                  margin=0.15, fit=False)


def l_columns(slide, s):
    blocks = s["blocks"]
    h = CH - (0.95 if s.get("summary") else 0)
    geo = split_w(CW, len(blocks), 0.2)
    sz = common_size([(w, h, b.get("title"), b.get("points")) for (dx, w), b in zip(geo, blocks)])
    for (dx, w), b in zip(geo, blocks):
        acc = b.get("accent")
        block(slide, CX + dx, CY, w, h, b.get("title"), b.get("points"), size=sz, fit=False,
              line=ACCENT if acc else None, title_color=ACCENT if acc else BLACK)
    if s.get("summary"):
        sh = rect(slide, CX, CY + h + 0.2, CW, 0.75, fill=WHITE, line=ACCENT, line_w=1.25)
        fill_text(sh.text_frame, [{"text": s["summary"], "bold": True, "color": ACCENT}],
                  CW, 0.75, size=13, anchor=MSO_ANCHOR.MIDDLE, margin=0.15)


def l_rows(slide, s):
    blocks = s["blocks"]
    n = len(blocks)
    gap = 0.15
    rh = (CH - gap * (n - 1)) / n
    lw = s.get("title_width", 3.6)
    for i, b in enumerate(blocks):
        y = CY + i * (rh + gap)
        t = rect(slide, CX, y, lw, rh, fill=GREY_BOX, line=ACCENT if b.get("accent") else None)
        fill_text(t.text_frame, [{"text": b["title"], "bold": True,
                                  "color": ACCENT if b.get("accent") else BLACK}], lw, rh, size=15,
                  anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
        textbox(slide, CX + lw + 0.2, y, CW - lw - 0.2, rh,
                [{"text": p, "bullet": True, "space_before": 4} for p in b.get("points", [])],
                size=14, anchor=MSO_ANCHOR.MIDDLE)
        if i < n - 1:
            line(slide, CX + lw + 0.2, y + rh + gap / 2, CX + CW, y + rh + gap / 2)


def l_bullets(slide, s):
    paras = []
    if s.get("block_title"):
        paras.append({"text": s["block_title"], "bold": True, "rel_size": 1.08})
    paras += [{"text": p, "bullet": True, "space_before": 8} for p in s["points"]]
    r = rect(slide, CX, CY, CW, CH)
    fill_text(r.text_frame, paras, CW, CH, size=16, margin=0.3)


def l_progression(slide, s):
    blocks = s["blocks"]
    n = len(blocks)
    ah = 0.5
    arrow = rect(slide, CX, CY, CW, ah, fill=ACCENT, kind=MSO_SHAPE.RIGHT_ARROW)
    if s.get("axis_label"):
        fill_text(arrow.text_frame, [{"text": s["axis_label"], "bold": True, "color": WHITE}],
                  CW, ah, size=12, anchor=MSO_ANCHOR.MIDDLE, margin=0.2, fit=False)
    top = CY + ah + 0.2
    full = CY + CH - top
    geo = split_w(CW, n, 0.2)
    hs = [full * (0.6 + 0.4 * (i / (n - 1) if n > 1 else 1)) for i in range(n)]
    sz = common_size([(w, h, b.get("title"), b.get("points")) for (dx, w), h, b in zip(geo, hs, blocks)])
    for i, ((dx, w), h, b) in enumerate(zip(geo, hs, blocks)):
        last = i == n - 1 and s.get("accent_last", True)
        block(slide, CX + dx, top + full - h, w, h, b.get("title"), b.get("points"), size=sz,
              fit=False, line=ACCENT if last else None, title_color=ACCENT if last else BLACK)


def l_chevrons(slide, s):
    steps = s["steps"]
    n = len(steps)
    ch_h = 0.75
    geo = split_w(CW, n, 0.06)
    # the chevron geometry reserves its points: the real text area is narrower
    lab = min(fit_size([{"text": e["title"], "bold": True}], max(0.6, w - 0.9) - 0.1, ch_h, 14, 9)
              for (dx, w), e in zip(geo, steps))
    sz = common_size([(w, CH - ch_h - 0.2, None, e.get("points")) for (dx, w), e in zip(geo, steps)])
    for i, ((dx, w), e) in enumerate(zip(geo, steps)):
        kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        c = rect(slide, CX + dx, CY, w + (0.12 if i < n - 1 else 0), ch_h,
                 fill=ACCENT if e.get("accent", True) else GREY_DARK, kind=kind)
        fill_text(c.text_frame, [{"text": e["title"], "bold": True, "color": WHITE,
                                  "align": PP_ALIGN.CENTER}], max(0.6, w - 0.9), ch_h, size=lab,
                  anchor=MSO_ANCHOR.MIDDLE, margin=0.05, fit=False)
        block(slide, CX + dx, CY + ch_h + 0.2, w, CH - ch_h - 0.2, None, e.get("points"), size=sz, fit=False)


def l_matrix(slide, s):
    ax, ay, q = s["x_axis"], s["y_axis"], s["quadrants"]
    side = s.get("side_blocks")
    gw = CW * (0.62 if side else 1.0)
    lx = CX + 1.0
    gx_w = gw - 1.0
    gy, gh = CY, CH - 0.75
    # horizontal box of width gh, rotated 270° around its centre (CX + 0.2, grid middle)
    lab = textbox(slide, CX + 0.2 - gh / 2, gy + gh / 2 - 0.2, gh, 0.4,
                  [{"text": ay["label"], "bold": True, "align": PP_ALIGN.CENTER}],
                  size=11, color=GREY_TEXT, fit=False, margin=0.02)
    lab.rotation = 270
    textbox(slide, CX + 0.35, gy, 0.6, 0.4, [ay.get("high", "High")], size=9, color=GREY_TEXT,
            fit=False, margin=0.02)
    textbox(slide, CX + 0.35, gy + gh - 0.4, 0.6, 0.4, [ay.get("low", "Low")], size=9,
            color=GREY_TEXT, fit=False, margin=0.02, anchor=MSO_ANCHOR.BOTTOM)
    line(slide, lx - 0.08, gy + gh, lx - 0.08, gy, color=GREY_DARK, width=1.25, arrow=True)
    line(slide, lx - 0.08, gy + gh + 0.08, lx + gx_w, gy + gh + 0.08, color=GREY_DARK, width=1.25, arrow=True)
    textbox(slide, lx, gy + gh + 0.15, gx_w, 0.35, [{"text": ax["label"], "bold": True,
            "align": PP_ALIGN.CENTER}], size=11, color=GREY_TEXT, fit=False, margin=0.02)
    textbox(slide, lx, gy + gh + 0.15, 1.5, 0.35, [ax.get("low", "Low")], size=9,
            color=GREY_TEXT, fit=False, margin=0.02)
    textbox(slide, lx + gx_w - 1.5, gy + gh + 0.15, 1.5, 0.35, [{"text": ax.get("high", "High"),
            "align": PP_ALIGN.RIGHT}], size=9, color=GREY_TEXT, fit=False, margin=0.02)
    qw, qh = (gx_w - 0.1) / 2, (gh - 0.1) / 2
    pos = [(0, 0), (1, 0), (0, 1), (1, 1)]  # top-left, top-right, bottom-left, bottom-right
    sz = common_size([(qw, qh, quad.get("title"), quad.get("points")) for quad in q], size=13)
    for (cx, cy), quad in zip(pos, q):
        acc = quad.get("accent")
        block(slide, lx + cx * (qw + 0.1), gy + cy * (qh + 0.1), qw, qh, quad.get("title"),
              quad.get("points"), line=ACCENT if acc else None,
              title_color=ACCENT if acc else BLACK, size=sz, fit=False)
    if side:
        side_blocks(slide, side, CX + gw + 0.3, CW - gw - 0.3)


def _cell_border_bottom(cell, color="BFBFBF", w=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag))
        if tag == "a:lnB":
            ln.set("w", str(w))
            etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", color)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))


def harvey_ball(slide, cx, cy, quarters, color, d=0.28):
    """Harvey ball drawn as shapes (0 to 4 quarters): same rendering whatever the font."""
    x, y = cx - d / 2, cy - d / 2
    if quarters >= 4:
        rect(slide, x, y, d, d, fill=color, kind=MSO_SHAPE.OVAL)
        return
    rect(slide, x, y, d, d, fill=WHITE, line=color, kind=MSO_SHAPE.OVAL, line_w=1.25)
    if quarters > 0:
        pie = rect(slide, x, y, d, d, fill=color, kind=MSO_SHAPE.PIE)
        # angles in degrees × 0.6 (python-pptx units); start at 12 o'clock, clockwise
        pie.adjustments[0] = 270 * 0.6
        pie.adjustments[1] = ((270 + 90 * quarters) % 360) * 0.6


def l_table(slide, s):
    cols, rows = s["columns"], s["rows"]
    acc = s.get("accent_column")
    nr, nc = len(rows) + 1, len(cols)
    rh = min(0.7, CH / nr)
    gf = slide.shapes.add_table(nr, nc, Inches(CX), Inches(CY), Inches(CW), Inches(rh * nr))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    etree.SubElement(tblPr, qn("a:tableStyleId")).text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    first_w = s.get("first_column_width", CW * 0.26)
    other = (CW - first_w) / (nc - 1) if nc > 1 else CW
    tbl.columns[0].width = Inches(first_w)
    for j in range(1, nc):
        tbl.columns[j].width = Inches(other)
    size = 13 if nr <= 7 else 11
    balls = []
    for i in range(nr):
        tbl.rows[i].height = Inches(rh)
        for j in range(nc):
            cell = tbl.cell(i, j)
            raw = str(cols[j] if i == 0 else (rows[i - 1][j] if j < len(rows[i - 1]) else ""))
            harvey = raw.strip() in HARVEY
            if harvey:
                cx = CX + (first_w + (j - 1) * other + other / 2 if j else first_w / 2)
                balls.append((cx, CY + i * rh + rh / 2, HARVEY[raw.strip()], j == acc))
            cell.fill.solid()
            cell.fill.fore_color.rgb = GREY_BOX if i == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.1)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if (j > 0 and (i == 0 or harvey)) else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = "" if harvey else raw
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = i == 0 or j == 0
            r.font.color.rgb = ACCENT if (acc is not None and j == acc and i == 0) else BLACK
            _cell_border_bottom(cell, "1A1A1A" if i == 0 else "BFBFBF")
    for cx, cy, v, is_acc in balls:
        harvey_ball(slide, cx, cy, v, ACCENT if is_acc else BLACK)
    if acc is not None:
        rect(slide, CX + first_w + (acc - 1) * other, CY, other, rh * nr, fill=None, line=ACCENT, line_w=1.5)
    if s.get("legend"):
        textbox(slide, CX, CY + rh * nr + 0.1, CW, 0.35, [s["legend"]], size=9,
                color=GREY_TEXT, fit=False, margin=0.02)


def _fmt(fmt, unit):
    return (fmt or "0") + (f'"{unit}"' if unit else "")


def _chart(slide, g, x, y, w, h):
    kind = {"bars": XL_CHART_TYPE.COLUMN_CLUSTERED, "hbars": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS, "stacked": XL_CHART_TYPE.COLUMN_STACKED}[g.get("type", "bars")]
    horizontal = kind == XL_CHART_TYPE.BAR_CLUSTERED
    cd = CategoryChartData()
    cats = g["categories"]
    cd.categories = list(reversed(cats)) if horizontal else cats  # first category on top
    for se in g["series"]:
        cd.add_series(se.get("name", ""), list(reversed(se["values"])) if horizontal else se["values"])
    ch = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_title = False
    ch.font.name = FONT
    ch.font.size = Pt(12)
    ch.font.color.rgb = BLACK
    multi = len(g["series"]) > 1
    ch.has_legend = multi
    if multi:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
    plot = ch.plots[0]
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    if kind != XL_CHART_TYPE.LINE_MARKERS:
        plot.gap_width = 60
        va.minimum_scale = 0  # bars start at zero
    if kind == XL_CHART_TYPE.COLUMN_STACKED:
        plot.overlap = 100
    ca = ch.category_axis
    ca.format.line.color.rgb = GREY_LINE
    ca.major_tick_mark = XL_TICK_MARK.NONE
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = _fmt(g.get("format"), g.get("unit"))
    dl.number_format_is_linked = False
    dl.font.size = Pt(11)
    dl.position = {XL_CHART_TYPE.LINE_MARKERS: XL_LABEL_POSITION.ABOVE,
                   XL_CHART_TYPE.COLUMN_STACKED: XL_LABEL_POSITION.CENTER}.get(kind, XL_LABEL_POSITION.OUTSIDE_END)
    palette = [ACCENT, GREY_DARK, GREY_LINE, RGBColor(0x40, 0x40, 0x40)]
    accents = set(g.get("accent", []))
    n = len(cats)
    for k, series in enumerate(plot.series):
        if kind == XL_CHART_TYPE.LINE_MARKERS:
            series.format.line.color.rgb = palette[k % 4]
            series.format.line.width = Pt(2.25)
            series.smooth = False
            continue
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = palette[k % 4]
        if kind == XL_CHART_TYPE.COLUMN_STACKED and k == 0:
            series.data_labels.font.color.rgb = WHITE
        if not multi:
            for i in range(n):
                pt = series.points[(n - 1 - i) if horizontal else i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = ACCENT if (i in accents or not accents) else GREY_MID
    return ch


def _chart_title(slide, txt, x, y, w):
    textbox(slide, x, y, w, 0.6, [{"text": txt, "bold": True}], size=13, min_size=11, margin=0.02)


def l_chart_text(slide, s):
    g = s["chart"]
    gw = CW * s.get("chart_share", 0.6)
    top = CY
    if g.get("title"):
        _chart_title(slide, g["title"], CX, CY, gw)
        top = CY + 0.65
    _chart(slide, g, CX, top, gw, CY + CH - top)
    side_blocks(slide, s["blocks"], CX + gw + 0.3, CW - gw - 0.3)


def l_big_number(slide, s):
    blocks = s.get("blocks")
    lw = CW * (0.45 if blocks else 1.0)
    textbox(slide, CX, CY + 0.6, lw, 2.2, [{"text": s["number"], "bold": True, "color": ACCENT,
            "align": PP_ALIGN.CENTER}], size=80, min_size=40, anchor=MSO_ANCHOR.BOTTOM, margin=0.05)
    textbox(slide, CX + 0.3, CY + 2.9, lw - 0.6, 1.5, [{"text": s["label"], "align": PP_ALIGN.CENTER}],
            size=18, min_size=12, color=GREY_TEXT, margin=0.05)
    if blocks:
        side_blocks(slide, blocks, CX + lw + 0.3, CW - lw - 0.3)


def _num(v, fmt, signed=False):
    dec = len(fmt.split(".")[1]) if fmt and "." in fmt else 0
    t = f"{abs(v):,.{dec}f}"
    sign = ("+" if v >= 0 else "−") if signed else ("−" if v < 0 else "")
    return sign + t


def l_waterfall(slide, s):
    steps = s["steps"]
    blocks = s.get("blocks")
    gw = CW * (0.62 if blocks else 1.0)
    top = CY
    if s.get("chart_title"):
        _chart_title(slide, s["chart_title"], CX, CY, gw)
        top = CY + 0.65
    base, val, kinds, run = [], [], [], 0.0
    for e in steps:
        v = float(e["value"])
        if e.get("total"):
            base.append(0); val.append(v); run = v; kinds.append("total")
        elif v >= 0:
            base.append(run); val.append(v); run += v; kinds.append("up")
        else:
            base.append(run + v); val.append(-v); run += v; kinds.append("down")
    cd = CategoryChartData()
    cd.categories = [e["label"] for e in steps]
    cd.add_series("base", base)
    cd.add_series("value", val)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(CX), Inches(top), Inches(gw),
                                Inches(CY + CH - top), cd).chart
    ch.has_title = False
    ch.font.name = FONT
    ch.font.size = Pt(12)
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = 50
    plot.overlap = 100
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.minimum_scale = 0
    ch.category_axis.format.line.color.rgb = GREY_LINE
    ch.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    sb, sv = plot.series
    sb.format.fill.background()
    sb.format.line.fill.background()
    fmt, unit = s.get("format", "0"), s.get("unit", "")
    for i, k in enumerate(kinds):
        pt = sv.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = {"total": ACCENT, "up": GREY_DARK, "down": GREY_LINE}[k]
        v = val[i] if k != "down" else -val[i]
        dlab = pt.data_label
        dlab.has_text_frame = True
        dlab.text_frame.text = _num(v, fmt, signed=(k != "total")) + unit
        dlab.position = XL_LABEL_POSITION.INSIDE_END
        r = dlab.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(11)
        r.font.bold = k == "total"
        r.font.color.rgb = WHITE if k in ("total", "up") else BLACK
    if blocks:
        side_blocks(slide, blocks, CX + gw + 0.3, CW - gw - 0.3)


def l_funnel(slide, s):
    stages = s["stages"]
    blocks = s.get("blocks")
    fw = CW * (0.6 if blocks else 1.0)
    n = len(stages)
    gap = 0.1
    sh = (CH - gap * (n - 1)) / n
    acc = s.get("accent_stage")
    for i, st in enumerate(stages):
        w = fw * (1.0 - 0.55 * i / max(1, n - 1))
        x = CX + (fw - w) / 2
        y = CY + i * (sh + gap)
        r = rect(slide, x, y, w, sh, fill=ACCENT if i == acc else GREY_BOX)
        col = WHITE if i == acc else BLACK
        txt = st["label"] + (f"  {st['value']}" if st.get("value") else "")
        fill_text(r.text_frame, [{"text": txt, "bold": True, "color": col, "align": PP_ALIGN.CENTER}],
                  w, sh, size=14, anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
    if blocks:
        side_blocks(slide, blocks, CX + fw + 0.3, CW - fw - 0.3)


def l_quotes(slide, s):
    qs = s["quotes"]
    qw = CW * 0.55
    n = len(qs)
    gap = 0.15
    qh = (CH - gap * (n - 1)) / n
    paras_all = [[{"text": f"“{q['text']}”", "italic": True},
                  {"text": q.get("author", ""), "color": GREY_TEXT, "rel_size": 0.85, "space_before": 6}]
                 for q in qs]
    sz = min(fit_size([_para(p) for p in paras if p["text"]], qw - 0.5, qh - 0.25, 15, 10) for paras in paras_all)
    for i, paras in enumerate(paras_all):
        y = CY + i * (qh + gap)
        rect(slide, CX, y, 0.08, qh, fill=ACCENT)
        r = rect(slide, CX + 0.08, y, qw - 0.08, qh, fill=GREY_BOX)
        fill_text(r.text_frame, paras, qw - 0.08, qh, size=sz, anchor=MSO_ANCHOR.MIDDLE, margin=0.2, fit=False)
    side_blocks(slide, s["blocks"], CX + qw + 0.3, CW - qw - 0.3)


def l_before_after(slide, s):
    aw = 1.0
    w = (CW - aw - 0.4) / 2
    a, b = s["before"], s["after"]
    sz = common_size([(w, CH, a.get("title"), a.get("points")), (w, CH, b.get("title"), b.get("points"))])
    block(slide, CX, CY, w, CH, a.get("title"), a.get("points"), size=sz, fit=False)
    block(slide, CX + w + aw + 0.4, CY, w, CH, b.get("title"), b.get("points"), line=ACCENT,
          title_color=ACCENT, size=sz, fit=False)
    rect(slide, CX + w + 0.2, CY + CH / 2 - 0.45, aw, 0.9, fill=ACCENT, kind=MSO_SHAPE.RIGHT_ARROW)
    if s.get("transition"):
        textbox(slide, CX + w + 0.05, CY + CH / 2 + 0.55, aw + 0.3, 1.2,
                [{"text": s["transition"], "align": PP_ALIGN.CENTER}], size=10, color=GREY_TEXT, margin=0.02)


def l_timeline(slide, s):
    ph = s["phases"]
    n = len(ph)
    ly = CY + 0.75
    line(slide, CX, ly, CX + CW, ly, color=ACCENT, width=2.5, arrow=True)
    geo = split_w(CW, n, 0.2)
    bh = CY + CH - ly - 0.3
    sz = common_size([(w, bh, p.get("title"), p.get("points")) for (dx, w), p in zip(geo, ph)])
    for (dx, w), p in zip(geo, ph):
        cx = CX + dx + w / 2
        textbox(slide, CX + dx, CY, w, 0.55, [{"text": p["period"], "bold": True,
                "align": PP_ALIGN.CENTER, "color": ACCENT}], size=12, margin=0.02, anchor=MSO_ANCHOR.BOTTOM)
        rect(slide, cx - 0.11, ly - 0.11, 0.22, 0.22, fill=ACCENT, kind=MSO_SHAPE.OVAL)
        block(slide, CX + dx, ly + 0.3, w, bh, p.get("title"), p.get("points"), size=sz, fit=False)


def l_gantt(slide, s):
    periods, tasks = s["periods"], s["tasks"]
    np_, nt = len(periods), len(tasks)
    lw = s.get("label_width", 3.4)
    gx, gw = CX + lw, CW - lw
    pw = gw / np_
    hh = 0.45
    rh = min(0.6, (CH - hh) / nt)
    for k, p in enumerate(periods):
        r = rect(slide, gx + k * pw, CY, pw - 0.03, hh, fill=GREY_BOX)
        fill_text(r.text_frame, [{"text": p, "bold": True, "align": PP_ALIGN.CENTER}], pw, hh, size=11,
                  anchor=MSO_ANCHOR.MIDDLE, margin=0.03)
    lab_sz = min(fit_size([{"text": t["label"]}], lw - 0.2, rh, 13, 9, quiet=True) for t in tasks)
    for i, t in enumerate(tasks):
        y = CY + hh + i * rh
        textbox(slide, CX, y, lw - 0.1, rh, [t["label"]], size=lab_sz, anchor=MSO_ANCHOR.MIDDLE,
                fit=False, margin=0.05)
        line(slide, CX, y + rh, CX + CW, y + rh, color=GREY_BOX)
        start, end = int(t["start"]), int(t.get("end", t["start"]))
        if t.get("milestone"):
            d = rh * 0.5
            rect(slide, gx + (end + 0.5) * pw - d / 2, y + (rh - d) / 2, d, d,
                 fill=ACCENT, kind=MSO_SHAPE.DIAMOND)
        else:
            rect(slide, gx + start * pw + 0.05, y + rh * 0.25, (end - start + 1) * pw - 0.1, rh * 0.5,
                 fill=ACCENT if t.get("accent") else GREY_DARK)


def l_tree(slide, s):
    br = s["branches"]
    n = len(br)
    rw, bw = 2.6, 3.2
    gap = 0.2
    bh = (CH - gap * (n - 1)) / n
    root = rect(slide, CX, CY + CH / 2 - 0.9, rw, 1.8, fill=ACCENT)
    fill_text(root.text_frame, [{"text": s["root"], "bold": True, "color": WHITE}], rw, 1.8,
              size=15, anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
    bx = CX + rw + 0.6
    fx = bx + bw + 0.5
    fw = CX + CW - fx
    xm = CX + rw + 0.3  # vertical spine that distributes the branches
    line(slide, CX + rw, CY + CH / 2, xm, CY + CH / 2, color=GREY_DARK, width=1.25)
    line(slide, xm, CY + bh / 2, xm, CY + (n - 1) * (bh + gap) + bh / 2, color=GREY_DARK, width=1.25)
    for i, b in enumerate(br):
        y = CY + i * (bh + gap)
        line(slide, xm, y + bh / 2, bx, y + bh / 2, color=GREY_DARK, width=1.25)
        r = rect(slide, bx, y, bw, bh, fill=GREY_BOX)
        fill_text(r.text_frame, [{"text": b["text"], "bold": True}], bw, bh, size=14,
                  anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
        if b.get("leaves"):
            line(slide, bx + bw, y + bh / 2, fx, y + bh / 2, color=GREY_DARK, width=1.25)
            textbox(slide, fx, y, fw, bh, [{"text": f, "bullet": True, "space_before": 3}
                                           for f in b["leaves"]], size=13, anchor=MSO_ANCHOR.MIDDLE)


def l_recommendation(slide, s):
    rh = 1.25
    r = rect(slide, CX, CY, CW, rh, fill=ACCENT)
    fill_text(r.text_frame, [{"text": s["recommendation"], "bold": True, "color": WHITE}], CW, rh,
              size=16, anchor=MSO_ANCHOR.MIDDLE, margin=0.25)
    pil = s["pillars"]
    top = CY + rh + 0.3
    geo = split_w(CW, len(pil), 0.2)
    ph = CY + CH - top
    sz = common_size([(w, ph, p.get("title"), p.get("points")) for (dx, w), p in zip(geo, pil)])
    for (dx, w), p in zip(geo, pil):
        line(slide, CX + dx + w / 2, CY + rh, CX + dx + w / 2, top, color=ACCENT, width=1.5)
        block(slide, CX + dx, top, w, ph, p.get("title"), p.get("points"), size=sz, fit=False)



# ---------------------------------------------------------------- benchmark layouts
def _cell_value(c):
    """A cell is a string, "[0]".."[4]" for a Harvey ball, or {"text", "calc"}; calc values print in blue."""
    if isinstance(c, dict):
        return str(c.get("text", "")), bool(c.get("calc"))
    return str(c), False


def _squarify(vals, x, y, w, h):
    """Squarified treemap (Bruls et al.): vals sorted descending, sum(vals) == w * h."""
    def worst(row, side):
        s = sum(row)
        return max(max(side * side * r / (s * s), s * s / (side * side * r)) for r in row)
    out, vals = [], list(vals)
    while vals:
        side = min(w, h)
        row, i = [vals[0]], 1
        while i < len(vals) and worst(row + [vals[i]], side) <= worst(row, side):
            row.append(vals[i])
            i += 1
        s = sum(row)
        if w >= h:
            rw, yy = s / h, y
            for r in row:
                out.append((x, yy, rw, r / rw))
                yy += r / rw
            x, w = x + rw, w - rw
        else:
            rh, xx = s / w, x
            for r in row:
                out.append((xx, y, r / rh, rh))
                xx += r / rh
            y, h = y + rh, h - rh
        vals = vals[i:]
    return out


def l_market_map(slide, s):
    items = [it for it in s["items"] if float(it.get("value", 0)) > 0]
    blocks = s.get("blocks")
    mw = CW * (s.get("map_share", 0.68) if blocks else 1.0)
    lh = 0.35
    mh = CH - lh - 0.1
    groups = []
    for it in items:
        if it.get("group", "") not in groups:
            groups.append(it.get("group", ""))
    if len(groups) > len(GROUP_FILLS):
        warn(f"{len(groups)} groups; {len(GROUP_FILLS)} colours at most, merge small groups")
    totals = {g: sum(float(i["value"]) for i in items if i.get("group", "") == g) for g in groups}
    order = sorted(groups, key=lambda g: -totals[g])
    total = sum(totals.values())
    unit = s.get("unit", "")
    grects = _squarify([totals[g] * mw * mh / total for g in order], CX, CY, mw, mh)
    for g, (gx, gy, gw, gh) in zip(order, grects):
        k = groups.index(g) % len(GROUP_FILLS)
        members = sorted([i for i in items if i.get("group", "") == g], key=lambda i: -float(i["value"]))
        pad = 0.03
        ix, iy, iw, ih = gx + pad, gy + pad, max(0.01, gw - 2 * pad), max(0.01, gh - 2 * pad)
        sub = _squarify([float(m["value"]) * iw * ih / totals[g] for m in members], ix, iy, iw, ih)
        for m, (rx, ry, rw, rh) in zip(members, sub):
            r = rect(slide, rx + 0.015, ry + 0.015, max(0.01, rw - 0.03), max(0.01, rh - 0.03),
                     fill=GROUP_FILLS[k])
            if rw > 0.55 and rh > 0.32:
                col = BLACK if k in GROUP_TEXT_DARK else WHITE
                paras = [{"text": m["label"], "bold": True, "color": col},
                         {"text": f"{m['value']}{unit}", "color": col, "rel_size": 0.9, "space_before": 1}]
                sz = fit_size([_para(p) for p in paras], rw - 0.12, rh - 0.1, 12, 7, quiet=True)
                fill_text(r.text_frame, paras, rw, rh, size=sz, margin=0.05, fit=False)
    # legend
    x = CX
    for k, g in enumerate(groups):
        if not g:
            continue
        rect(slide, x, CY + CH - lh + 0.1, 0.18, 0.18, fill=GROUP_FILLS[k % len(GROUP_FILLS)])
        w = min(2.6, 0.3 + len(g) * 0.075)
        textbox(slide, x + 0.22, CY + CH - lh + 0.02, w, 0.34, [g], size=10, color=GREY_TEXT,
                fit=False, margin=0.02, anchor=MSO_ANCHOR.MIDDLE)
        x += 0.25 + w
    if blocks:
        side_blocks(slide, blocks, CX + mw + 0.3, CW - mw - 0.3)


def l_band_matrix(slide, s):
    bands, rows = s["column_bands"], s["rows"]
    cols = [c for b in bands for c in b["columns"]]
    nc = len(cols) + 1
    data = []  # ("group", label) or ("row", row)
    last = None
    for r in rows:
        g = r.get("group")
        if g and g != last:
            data.append(("group", g))
            last = g
        data.append(("row", r))
    nr = 2 + len(data)
    avail = CH - (0.35 if s.get("legend") else 0)
    rh = min(0.42, avail / nr)
    size = max(7.0, min(11.0, rh * 72 * 0.55))
    first_w = s.get("first_column_width", 2.6)
    other = (CW - first_w) / len(cols)
    gf = slide.shapes.add_table(nr, nc, Inches(CX), Inches(CY), Inches(CW), Inches(rh * nr))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    etree.SubElement(tblPr, qn("a:tableStyleId")).text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    tbl.columns[0].width = Inches(first_w)
    for j in range(1, nc):
        tbl.columns[j].width = Inches(other)
    for i in range(nr):
        tbl.rows[i].height = Inches(rh)

    def put(cell, text, bold=False, color=BLACK, align=PP_ALIGN.CENTER, fill=WHITE, border="BFBFBF", sz=size):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.05)
        cell.margin_top = cell.margin_bottom = Inches(0.01)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.color.rgb = color
        _cell_border_bottom(cell, border)

    # band header row (merged cells) and column header row
    put(tbl.cell(0, 0), s.get("band_header", ""), bold=True, align=PP_ALIGN.LEFT)
    put(tbl.cell(1, 0), s.get("row_header", ""), bold=True, align=PP_ALIGN.LEFT, fill=GREY_BOX, border="1A1A1A")
    j = 1
    for k, b in enumerate(bands):
        n = len(b["columns"])
        fill = ACCENT if k % 2 == 0 else GREY_DARK
        put(tbl.cell(0, j), b["label"], bold=True, color=WHITE, fill=fill)
        if n > 1:
            tbl.cell(0, j).merge(tbl.cell(0, j + n - 1))
        for c in b["columns"]:
            put(tbl.cell(1, j), c, bold=True, fill=GREY_BOX, border="1A1A1A", sz=max(7.0, size - 0.5))
            j += 1
    balls = []
    for i, (kind, val) in enumerate(data, start=2):
        if kind == "group":
            put(tbl.cell(i, 0), val, bold=True, color=ACCENT, align=PP_ALIGN.LEFT, border="1A1A1A")
            for jj in range(1, nc):
                put(tbl.cell(i, jj), "", border="1A1A1A")
            tbl.cell(i, 0).merge(tbl.cell(i, nc - 1))
            continue
        acc = val.get("accent")
        put(tbl.cell(i, 0), val["label"], bold=True, color=ACCENT if acc else BLACK, align=PP_ALIGN.LEFT)
        cells = val.get("cells", [])
        for jj in range(1, nc):
            raw = cells[jj - 1] if jj - 1 < len(cells) else ""
            text, calc = _cell_value(raw)
            if text.strip() in HARVEY:
                balls.append((CX + first_w + (jj - 1) * other + other / 2, CY + i * rh + rh / 2, HARVEY[text.strip()]))
                text = ""
            sym = text.strip() in ("●", "○", "◐")
            put(tbl.cell(i, jj), text, color=CALC_BLUE if calc else BLACK, sz=size + 3 if sym else size)
    d = min(0.22, rh * 0.7)
    for cx, cy, v in balls:
        harvey_ball(slide, cx, cy, v, BLACK, d=d)
    legend = s.get("legend")
    if legend or any(_cell_value(c)[1] for r in rows for c in r.get("cells", [])):
        txt = (legend + " · " if legend else "") + "Blue: computed by the team"
        textbox(slide, CX, CY + rh * nr + 0.08, CW, 0.3, [txt], size=9, color=GREY_TEXT, fit=False, margin=0.02)


def l_fact_sheet(slide, s):
    facts = s["facts"]
    blocks = s.get("blocks")
    fw = CW * (0.42 if blocks else 1.0)
    n = len(facts)
    rh = min(0.62, CH / n)
    lw = fw * 0.5
    sz = 13 if rh >= 0.5 else 11
    for i, f in enumerate(facts):
        y = CY + i * rh
        textbox(slide, CX, y, lw, rh, [{"text": f["label"], "color": GREY_TEXT}], size=sz,
                anchor=MSO_ANCHOR.MIDDLE, margin=0.05, min_size=8)
        text, calc = _cell_value(f.get("value", ""))
        textbox(slide, CX + lw, y, fw - lw, rh, [{"text": text, "bold": True,
                "color": CALC_BLUE if calc else BLACK}], size=sz, anchor=MSO_ANCHOR.MIDDLE, margin=0.05, min_size=8)
        line(slide, CX, y + rh, CX + fw, y + rh, color=GREY_LINE, width=0.75)
    if any(_cell_value(f.get("value", ""))[1] for f in facts):
        textbox(slide, CX, CY + n * rh + 0.05, fw, 0.3, ["Blue: computed by the team"], size=9,
                color=GREY_TEXT, fit=False, margin=0.02)
    if blocks:
        side_blocks(slide, blocks, CX + fw + 0.3, CW - fw - 0.3)


BUILDERS = {k: globals()[f"l_{k}"] for k in LAYOUTS}


# ---------------------------------------------------------------- main
def build(story, out):
    global ACCENT, FONT
    meta = story.get("meta", {})
    if meta.get("accent"):
        ACCENT = RGBColor.from_string(meta["accent"].lstrip("#"))
    FONT = meta.get("font", FONT)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]
    for i, s in enumerate(story["slides"], start=1):
        _CTX["slide"] = i
        lay = s.get("layout")
        if lay not in BUILDERS:
            raise SystemExit(f"slide {i}: unknown layout '{lay}' (see layouts_spec.py)")
        missing = [f for f in LAYOUTS[lay][0] if f not in s]
        if missing:
            raise SystemExit(f"slide {i} ({lay}): missing fields {missing}")
        slide = prs.slides.add_slide(blank)
        if lay == "cover":
            l_cover(slide, s)
        else:
            header(slide, s, i)
            BUILDERS[lay](slide, s)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("storyboard")
    ap.add_argument("-o", "--out", default="deck.pptx")
    a = ap.parse_args()
    story = json.loads(Path(a.storyboard).read_text(encoding="utf-8"))
    build(story, a.out)
    for w in WARNINGS:
        print("WARNING", w, file=sys.stderr)
    print(f"OK {a.out} ({len(story['slides'])} slides, {len(WARNINGS)} warning(s))")


if __name__ == "__main__":
    main()
